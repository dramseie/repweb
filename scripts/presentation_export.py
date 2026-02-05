#!/usr/bin/env python3
import json
import re
import sys
from datetime import datetime
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
except ImportError as exc:
    raise SystemExit("python-pptx is required. Install with: pip install python-pptx") from exc


def get_country_blocks(section):
    return {item.get("country"): item for item in section.get("items", [])}


def add_table(slide, left, top, width, rows):
    cols = ["Site ID", "Site Name", "Start", "End", "Confidence", "Status"]
    row_count = max(2, len(rows) + 1)
    table_shape = slide.shapes.add_table(row_count, len(cols), left, top, width, Inches(1.1))
    table = table_shape.table

    for idx, col in enumerate(cols):
        cell = table.cell(0, idx)
        cell.text = col
        for paragraph in cell.text_frame.paragraphs:
            paragraph.font.bold = True
            paragraph.font.size = Pt(10)

    for r_idx, row in enumerate(rows, start=1):
        values = [
            row.get("siteId") or "",
            row.get("siteName") or "",
            short_date(row.get("startDate")),
            short_date(row.get("endDate")),
            row.get("confidence") or "",
            row.get("status") or "",
        ]
        for c_idx, value in enumerate(values):
            cell = table.cell(r_idx, c_idx)
            cell.text = str(value)
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.size = Pt(9)

    return table_shape


def add_section(slide, y, title, meta, current_rows, next_rows):
    title_box = slide.shapes.add_textbox(Inches(0.6), y, Inches(12.0), Inches(0.3))
    title_tf = title_box.text_frame
    title_tf.text = title
    title_tf.paragraphs[0].font.bold = True
    title_tf.paragraphs[0].font.size = Pt(14)

    month_left = meta.get("currentMonth", "Current")
    month_right = meta.get("nextMonth", "Next")

    left_label = slide.shapes.add_textbox(Inches(0.6), y + Inches(0.35), Inches(5.6), Inches(0.2))
    left_label.text_frame.text = month_left
    left_label.text_frame.paragraphs[0].font.size = Pt(10)

    right_label = slide.shapes.add_textbox(Inches(6.3), y + Inches(0.35), Inches(5.6), Inches(0.2))
    right_label.text_frame.text = month_right
    right_label.text_frame.paragraphs[0].font.size = Pt(10)

    if current_rows:
        add_table(slide, Inches(0.6), y + Inches(0.6), Inches(5.6), current_rows)
    else:
        empty_left = slide.shapes.add_textbox(Inches(0.6), y + Inches(0.6), Inches(5.6), Inches(0.3))
        empty_left.text_frame.text = "No planned assessments."
        empty_left.text_frame.paragraphs[0].font.size = Pt(9)

    if next_rows:
        add_table(slide, Inches(6.3), y + Inches(0.6), Inches(5.6), next_rows)
    else:
        empty_right = slide.shapes.add_textbox(Inches(6.3), y + Inches(0.6), Inches(5.6), Inches(0.3))
        empty_right.text_frame.text = "No planned assessments."
        empty_right.text_frame.paragraphs[0].font.size = Pt(9)


def add_issue_table(slide, left, top, width, rows):
    cols = [
        "Store Name",
        "Store ID",
        "Description",
        "Priority (CHML)",
        "Responsible Party",
        "Action to be taken (DD.MM.YY - NS)",
        "Date to be resolved (DD.MM.YY)",
    ]
    row_count = max(2, len(rows) + 1)
    table_shape = slide.shapes.add_table(row_count, len(cols), left, top, width, Inches(3.6))
    table = table_shape.table

    for idx, col in enumerate(cols):
        cell = table.cell(0, idx)
        cell.text = col
        for paragraph in cell.text_frame.paragraphs:
            paragraph.font.bold = True
            paragraph.font.size = Pt(9)

    for r_idx, row in enumerate(rows, start=1):
        values = [
            row.get("storeName") or "",
            row.get("storeId") or "",
            row.get("description") or "",
            row.get("priority") or "",
            row.get("responsibleParty") or "",
            short_date(row.get("actionRequired")),
            short_date(row.get("resolveDate")),
        ]
        for c_idx, value in enumerate(values):
            cell = table.cell(r_idx, c_idx)
            cell.text = str(value)
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.size = Pt(8)

    return table_shape


def strip_html(value):
    if not value:
        return ""
    text = re.sub(r"<[^>]+>", "", str(value))
    return re.sub(r"\s+", " ", text).strip()


def short_date(value):
    if value in (None, "", "—"):
        return ""
    if isinstance(value, (int, float)):
        return str(value)
    raw = str(value).strip()
    if not raw:
        return ""
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%d/%m/%Y", "%d.%m.%Y", "%d.%m.%y"):
        try:
            return datetime.strptime(raw[:10], fmt).strftime("%d.%m.%y")
        except ValueError:
            continue
    try:
        cleaned = raw.replace("Z", "")
        parsed = datetime.fromisoformat(cleaned)
        return parsed.strftime("%d.%m.%y")
    except ValueError:
        return raw


def add_simple_table(slide, title, headers, rows, left=Inches(0.6), top=Inches(1.2), width=Inches(12.0), height=Inches(5.0)):
    title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.3), Inches(12.0), Inches(0.4))
    title_tf = title_box.text_frame
    title_tf.text = title
    title_tf.paragraphs[0].font.size = Pt(20)
    title_tf.paragraphs[0].font.bold = True

    row_count = max(2, len(rows) + 1)
    table_shape = slide.shapes.add_table(row_count, len(headers), left, top, width, height)
    table = table_shape.table
    for idx, col in enumerate(headers):
        cell = table.cell(0, idx)
        cell.text = col
        for paragraph in cell.text_frame.paragraphs:
            paragraph.font.bold = True
            paragraph.font.size = Pt(10)
    for r_idx, row in enumerate(rows, start=1):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = str(value)
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.size = Pt(9)


def main(input_path, output_path):
    data = json.loads(Path(input_path).read_text(encoding="utf-8"))
    prs = Presentation()

    generated_at = data.get("generatedAt") or ""

    assessments = data.get("plannedAssessments", {})
    installations = data.get("plannedInstallations", {})
    post_deployment = data.get("postDeployment", {})
    issue_log = data.get("issueLog", {})
    overview_items = data.get("overviewItems", [])
    timeline = data.get("timeline", {})
    planned_week_rows = data.get("plannedWeekRows", [])
    highlights = strip_html(data.get("highlights", ""))
    trend_overrides = data.get("trendOverrides", {})
    overview_overrides = data.get("overviewOverrides", {})

    assessment_blocks = get_country_blocks(assessments)
    installation_blocks = get_country_blocks(installations)
    post_blocks = get_country_blocks(post_deployment)
    issue_blocks = get_country_blocks(issue_log)

    country_list = data.get("countries") or sorted(
        {
            c
            for c in list(assessment_blocks.keys())
            + list(installation_blocks.keys())
            + list(post_blocks.keys())
            + list(issue_blocks.keys())
            if c
        }
    )

    title_slide = prs.slides.add_slide(prs.slide_layouts[6])
    title_box = title_slide.shapes.add_textbox(Inches(0.6), Inches(2.2), Inches(12.0), Inches(1.0))
    title_tf = title_box.text_frame
    title_tf.text = "Presentation Export"
    title_tf.paragraphs[0].font.size = Pt(32)
    title_tf.paragraphs[0].font.bold = True
    if generated_at:
        subtitle = title_slide.shapes.add_textbox(Inches(0.6), Inches(3.1), Inches(12.0), Inches(0.4))
        subtitle.text_frame.text = f"Generated {short_date(generated_at)}"
        subtitle.text_frame.paragraphs[0].font.size = Pt(14)

    if highlights:
        highlight_slide = prs.slides.add_slide(prs.slide_layouts[6])
        title_box = highlight_slide.shapes.add_textbox(Inches(0.6), Inches(0.3), Inches(12.0), Inches(0.4))
        title_tf = title_box.text_frame
        title_tf.text = "Highlights"
        title_tf.paragraphs[0].font.size = Pt(20)
        title_tf.paragraphs[0].font.bold = True
        body = highlight_slide.shapes.add_textbox(Inches(0.6), Inches(1.0), Inches(12.0), Inches(5.5))
        body_tf = body.text_frame
        body_tf.text = highlights
        body_tf.paragraphs[0].font.size = Pt(14)

    if overview_items:
        overview_rows = []
        for row in overview_items:
            country = row.get("country") or ""
            override = overview_overrides.get(country, {}) if isinstance(overview_overrides, dict) else {}
            rag = (override.get("rag") or row.get("rag") or "").strip()
            comment = (override.get("comment") or row.get("comment") or "").strip()
            overview_rows.append([
                country,
                row.get("stores") or "",
                row.get("assessed") or "",
                row.get("ongoingInstallations") or "",
                row.get("storesInstalled") or "",
                row.get("storeSignoff") or "",
                rag,
                comment,
            ])
        add_simple_table(
            prs.slides.add_slide(prs.slide_layouts[6]),
            "Programme Overview Per Country",
            ["Country", "Stores", "Assessed", "Ongoing Installations", "Installed", "Sign-off", "RAG", "Comment"],
            overview_rows,
        )

    if planned_week_rows:
        status_rows = []
        for row in planned_week_rows:
            status_rows.append([
                row.get("country") or row.get("Country") or "",
                row.get("site_name") or row.get("siteName") or row.get("Site_Name") or "",
                row.get("site_id") or row.get("siteId") or row.get("Site_ID") or "",
                row.get("task_name") or row.get("taskName") or row.get("Task_Name") or "",
                short_date(row.get("start_date") or row.get("startDate") or row.get("Start_Date")),
                short_date(row.get("end_date") or row.get("endDate") or row.get("End_Date")),
                row.get("status") or row.get("Status") or "",
                row.get("comment") or row.get("Comment") or "",
            ])
        add_simple_table(
            prs.slides.add_slide(prs.slide_layouts[6]),
            "Status of assessments and installations to start/finish",
            ["Country", "Site Name", "Site ID", "Activity", "Start", "End", "Status", "Comment"],
            status_rows,
        )

    timeline_items = timeline.get("items", []) if isinstance(timeline, dict) else []
    if timeline_items:
        timeline_rows = []
        for row in timeline_items:
            timeline_rows.append([
                row.get("country") or "",
                short_date(row.get("startDate")),
                short_date(row.get("installEndDate") or row.get("endDate")),
                short_date(row.get("endDate")),
            ])
        add_simple_table(
            prs.slides.add_slide(prs.slide_layouts[6]),
            "Timeline",
            ["Country", "Start", "Install End", "End"],
            timeline_rows,
        )

    trend_groups = {"green": [], "amber": [], "red": []}
    for row in overview_items:
        country = row.get("country") or ""
        override = trend_overrides.get(country, {}) if isinstance(trend_overrides, dict) else {}
        rag = (override.get("rag") or row.get("rag") or "").strip().lower()
        comment = (override.get("comment") or row.get("comment") or "").strip()
        if rag in trend_groups:
            trend_groups[rag].append([
                country,
                row.get("stores") or "",
                row.get("assessed") or "",
                row.get("ongoingInstallations") or "",
                row.get("storesInstalled") or "",
                comment,
            ])

    for rag_label, rows in [("Green", trend_groups["green"]), ("Amber", trend_groups["amber"]), ("Red", trend_groups["red"])]:
        if rows:
            add_simple_table(
                prs.slides.add_slide(prs.slide_layouts[6]),
                f"Country Trend: {rag_label}",
                ["Country", "Total", "Assessments", "Ongoing Installation", "Stores Installed", "Comment"],
                rows,
            )

    if issue_log.get("items"):
        issue_rows = []
        for block in issue_log.get("items", []):
            country = block.get("country") or ""
            for issue in block.get("issues", []):
                issue_rows.append([
                    country,
                    issue.get("storeName") or "",
                    issue.get("storeId") or "",
                    issue.get("description") or "",
                    issue.get("priority") or "",
                    issue.get("responsibleParty") or "",
                    short_date(issue.get("actionRequired")),
                    short_date(issue.get("resolveDate")),
                ])
        if issue_rows:
            add_simple_table(
                prs.slides.add_slide(prs.slide_layouts[6]),
                "General Issues",
                ["Country", "Site Name", "Site ID", "Description", "Priority", "Responsible Party", "Action", "Resolve Date"],
                issue_rows,
            )

    for country in country_list:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        title = slide.shapes.add_textbox(Inches(0.6), Inches(0.3), Inches(12.0), Inches(0.5))
        title.text_frame.text = country
        title.text_frame.paragraphs[0].font.size = Pt(20)
        title.text_frame.paragraphs[0].font.bold = True

        assessment = assessment_blocks.get(country, {})
        installation = installation_blocks.get(country, {})
        post = post_blocks.get(country, {})

        add_section(
            slide,
            Inches(1.0),
            "Planned Assessments",
            assessments.get("meta", {}),
            assessment.get("current", []),
            assessment.get("next", []),
        )

        add_section(
            slide,
            Inches(3.1),
            "Planned Installations",
            installations.get("meta", {}),
            installation.get("current", []),
            installation.get("next", []),
        )

        add_section(
            slide,
            Inches(5.2),
            "Post-Deployment & Sign-off",
            post_deployment.get("meta", {}),
            post.get("current", []),
            post.get("next", []),
        )

        issues = issue_blocks.get(country, {})
        if issues.get("issues"):
            issue_slide = prs.slides.add_slide(prs.slide_layouts[6])
            title = issue_slide.shapes.add_textbox(Inches(0.6), Inches(0.3), Inches(12.0), Inches(0.5))
            title.text_frame.text = f"{country} Issues"
            title.text_frame.paragraphs[0].font.size = Pt(20)
            title.text_frame.paragraphs[0].font.bold = True

            add_issue_table(issue_slide, Inches(0.6), Inches(1.2), Inches(12.0), issues.get("issues", []))

    prs.save(output_path)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("Usage: presentation_export.py <input.json> <output.pptx>")
    main(sys.argv[1], sys.argv[2])
